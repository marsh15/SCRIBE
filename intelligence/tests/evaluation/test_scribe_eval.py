import json

import docx
import pytest

from scribe_eval.cli import main
from scribe_eval.corpus import load_document_text
from scribe_eval.normalize import normalize
from scribe_eval.splits import assign_splits


def test_normalize():
    assert normalize("  Hello\t\nWORLD  ") == "hello world"
    assert normalize("It’s “fine” – ok—yes") == "it's \"fine\" - ok-yes"
    assert normalize("ＡＢ C") == "ab c"  # NFKC fullwidth + nbsp
    assert normalize("") == ""


def make_case(cid, category="a", split="dev", answerable=True, evidence=None, **kw):
    if evidence is None:
        evidence = [{"doc": "a.md", "version": 1, "quote": "alpha beta", "page": None, "grade": 2}] if answerable else []
    return {"id": cid, "split": split, "category": category, "tags": [], "question": "q?",
            "answerable": answerable, "reference_answer": "r", "must_include": [],
            "evidence": evidence, "forbidden": [], **kw}


@pytest.fixture
def corpus(tmp_path):
    c = tmp_path / "corpus"
    c.mkdir()
    (c / "a.md").write_text("Alpha   Beta gamma.\nDelta “quoted” text.", encoding="utf-8")
    (c / "dup.txt").write_text("same line. same line.", encoding="utf-8")
    (c / "t.csv").write_bytes("name,city\nJosé,Zürich\n".encode("latin-1", errors="replace"))
    d = docx.Document()
    d.add_paragraph("Intro paragraph")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "k"
    t.rows[0].cells[1].text = "v"
    d.add_paragraph("After table")
    d.save(str(c / "w.docx"))
    (c / "scan.pdf").write_bytes(b"not really a pdf")
    (c / "scan.truth.txt").write_text("Truth text of scan", encoding="utf-8")
    return c


def write(ds, rows):
    ds.mkdir(exist_ok=True)
    (ds / "cases.jsonl").write_text("".join(json.dumps(r) + "\n" for r in rows), encoding="utf-8")


def balanced_rows(**over):
    # 10 cases in 2 categories, 40/20/40
    rows = []
    for cat in ("a", "b"):
        for i, sp in enumerate(["dev", "dev", "validation", "test", "test"]):
            rows.append(make_case(f"{cat}{i}", category=cat, split=sp))
    return rows


def run(tmp_path, corpus, rows, capsys):
    ds = tmp_path / "ds"
    write(ds, rows)
    code = main(["validate", str(ds), "--corpus", str(corpus)])
    cap = capsys.readouterr()
    return code, cap.out + cap.err


def test_valid_dataset(tmp_path, corpus, capsys):
    code, out = run(tmp_path, corpus, balanced_rows(), capsys)
    assert code == 0 and "OK" in out and "TOTAL" in out


def test_quote_normalized_match_curly(tmp_path, corpus, capsys):
    rows = balanced_rows()
    rows[0]["evidence"][0]["quote"] = 'delta "quoted" text'
    code, _ = run(tmp_path, corpus, rows, capsys)
    assert code == 0


def test_quote_missing(tmp_path, corpus, capsys):
    rows = balanced_rows()
    rows[0]["evidence"][0]["quote"] = "nonexistent"
    code, out = run(tmp_path, corpus, rows, capsys)
    assert code == 1 and "found 0 times" in out and "[a0]" in out


def test_quote_duplicated(tmp_path, corpus, capsys):
    rows = balanced_rows()
    rows[0]["evidence"] = [{"doc": "dup.txt", "version": 1, "quote": "same line", "page": None, "grade": 1}]
    code, out = run(tmp_path, corpus, rows, capsys)
    assert code == 1 and "found 2 times" in out


def test_missing_doc_and_forbidden(tmp_path, corpus, capsys):
    rows = balanced_rows()
    rows[0]["evidence"][0]["doc"] = "nope.md"
    rows[1]["forbidden"] = [{"doc": "nope2.md"}]
    code, out = run(tmp_path, corpus, rows, capsys)
    assert code == 1 and "document not found" in out and "forbidden doc not found" in out


def test_quote_too_long(tmp_path, corpus, capsys):
    rows = balanced_rows()
    rows[0]["evidence"][0]["quote"] = "x" * 201
    code, out = run(tmp_path, corpus, rows, capsys)
    assert code == 1 and "at most 200 characters" in out


def test_unanswerable_with_evidence_and_answerable_without(tmp_path, corpus, capsys):
    rows = balanced_rows()
    rows[0]["answerable"] = False
    rows[1]["evidence"] = []
    code, out = run(tmp_path, corpus, rows, capsys)
    assert code == 1
    assert "[a0] unanswerable case must have empty evidence" in out
    assert "[a1] answerable case has no evidence" in out


def test_duplicate_ids_and_schema(tmp_path, corpus, capsys):
    rows = balanced_rows()
    rows[1]["id"] = rows[0]["id"]
    rows[2]["split"] = "bogus"
    code, out = run(tmp_path, corpus, rows, capsys)
    assert code == 1 and "duplicate id" in out and "schema error" in out


def test_unbalanced_categories(tmp_path, corpus, capsys):
    rows = balanced_rows()
    for i in range(10):
        rows.append(make_case(f"c{i}", category="a", split=["dev", "dev", "validation", "test", "test"][i % 5]))
    code, out = run(tmp_path, corpus, rows, capsys)
    assert code == 1 and "more than 2x the smallest" in out


def test_bad_split_ratios(tmp_path, corpus, capsys):
    rows = balanced_rows()
    for r in rows:
        r["split"] = "dev"
    code, out = run(tmp_path, corpus, rows, capsys)
    assert code == 1 and "split 'dev' is 100.0%" in out
    assert "missing from split(s)" in out


def test_assign_splits_determinism_and_stratification(tmp_path, corpus, capsys):
    rows = [make_case(f"{cat}{i:02d}", category=cat, split="dev") for cat in "abc" for i in range(10)]
    ds = tmp_path / "ds"
    write(ds, rows)
    assert main(["assign-splits", str(ds), "--seed", "7"]) == 0
    first = (ds / "cases.jsonl").read_text()
    write(ds, list(reversed(rows)))  # input order must not matter
    main(["assign-splits", str(ds), "--seed", "7"])
    assert {json.loads(l)["id"]: json.loads(l)["split"] for l in first.splitlines()} == {
        json.loads(l)["id"]: json.loads(l)["split"] for l in (ds / "cases.jsonl").read_text().splitlines()
    }
    out = [json.loads(l) for l in first.splitlines()]
    for cat in "abc":
        splits = [r["split"] for r in out if r["category"] == cat]
        assert (splits.count("dev"), splits.count("validation"), splits.count("test")) == (4, 2, 4)
    assert all(r["question"] == "q?" for r in out)  # other fields preserved
    other = assign_splits([dict(r) for r in rows], seed=8)
    assert {r["id"]: r["split"] for r in other} != {r["id"]: r["split"] for r in out}


def test_assign_remainder_to_dev():
    rows = [{"id": f"x{i}", "category": "a", "split": "test"} for i in range(7)]
    out = assign_splits(rows, 1)
    counts = [sum(r["split"] == s for r in out) for s in ("dev", "validation", "test")]
    assert counts == [3, 1, 3]


def test_loaders(corpus):
    assert normalize(load_document_text(corpus / "w.docx")) == "intro paragraph k | v after table"
    assert "name,city" in load_document_text(corpus / "t.csv")
    assert load_document_text(corpus / "a.md").startswith("Alpha")


def test_scanned_truth_sidecar(tmp_path, corpus, capsys):
    assert load_document_text(corpus / "scan.pdf") == "Truth text of scan"
    rows = balanced_rows()
    rows[0]["evidence"] = [{"doc": "scan.pdf", "version": 1, "quote": "truth text", "page": 1, "grade": 2}]
    code, _ = run(tmp_path, corpus, rows, capsys)
    assert code == 0
