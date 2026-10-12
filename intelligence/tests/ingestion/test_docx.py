"""DOCX parser: golden-v1 evidence, heading paths, tables, order, merges, errors."""

from __future__ import annotations

import io
import json

import docx
import pytest
from _helpers import CORPUS, GOLDEN
from scribe_eval.normalize import normalize
from scribe_ingest import IngestError, parse

MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _doc(name: str):
    return parse((CORPUS / name).read_bytes(), MIME)


def _entries():
    seen, out = set(), []
    for line in GOLDEN.read_text().splitlines():
        case = json.loads(line)
        for ev in case["evidence"]:
            key = (ev["doc"], ev["quote"])
            if ev["doc"].endswith(".docx") and key not in seen:
                seen.add(key)
                out.append(pytest.param(*key, id=f"{case['id']}-{ev['doc']}"))
    return out


ENTRIES = _entries()


def test_entries_collected():
    assert len(ENTRIES) >= 25


@pytest.mark.parametrize("doc,quote", ENTRIES)
def test_quote_lands_in_one_block(doc, quote):
    parsed = _doc(doc)
    q = normalize(quote)
    hits = [b for b in parsed.blocks if not b.furniture and q in normalize(b.text)]
    assert hits, f"quote not in any single block of {doc}: {quote!r}"
    assert all(b.page is None and b.bbox is None for b in parsed.blocks)


def _build(fn) -> bytes:
    d = docx.Document()
    fn(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_msa_heading_paths():
    p = _doc("master-services-agreement.docx")
    by_text = {b.text: b for b in p.blocks}
    h = by_text["2.2 Rates"]
    assert h.type == "heading" and h.heading_path == ["Master Services Agreement", "2. Services", "2.2 Rates"]
    late = next(b for b in p.blocks if b.text.startswith("Late payments accrue"))
    assert late.heading_path == ["Master Services Agreement", "4. Fees and payment", "4.2 Late payment"]
    # level pop: an H2 after an H3 drops the H3
    assert by_text["3. Delivery"].heading_path == ["Master Services Agreement", "3. Delivery"]


def test_msa_table_header_and_order():
    p = _doc("master-services-agreement.docx")
    i = next(i for i, b in enumerate(p.blocks) if b.type == "table")
    t = p.blocks[i]
    lines = t.text.split("\n")
    assert lines[0].startswith("| ") and lines[0].endswith(" |")
    assert t.meta["header"] == [c.strip() for c in lines[0][2:-2].split(" | ")]
    assert len(t.meta["rows"]) == len(lines) - 1
    assert "| Fuel surcharge | percent of freight | 6.5% |" in t.text
    # between the intro paragraph and the next heading
    assert p.blocks[i - 1].text.startswith("The rates below apply")
    assert p.blocks[i + 1].text == "3. Delivery"
    assert t.heading_path[-1] == "2.2 Rates"


def test_document_order_synthetic():
    def fn(d):
        d.add_paragraph("before")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "h1", "h2"
        t.cell(1, 0).text, t.cell(1, 1).text = "a", "b"
        d.add_paragraph("after")

    p = parse(_build(fn), MIME)
    assert [b.type for b in p.blocks] == ["paragraph", "table", "paragraph"]
    assert p.blocks[1].text == "| h1 | h2 |\n| a | b |"
    assert p.blocks[1].meta["header"] == ["h1", "h2"] and p.blocks[1].meta["rows"] == [["a", "b"]]


def test_merged_cells_not_duplicated():
    def fn(d):
        t = d.add_table(rows=3, cols=3)
        t.cell(0, 0).merge(t.cell(0, 2)).text = "Title"
        for c, v in enumerate(["x", "y", "z"]):
            t.cell(1, c).text = v
        t.cell(2, 0).merge(t.cell(2, 1)).text = "wide"
        t.cell(2, 2).text = "end"

    p = parse(_build(fn), MIME)
    t = p.blocks[0]
    assert t.text.count("Title") == 1 and t.text.count("wide") == 1
    assert t.text.split("\n")[1] == "| x | y | z |"
    assert t.text.split("\n")[2] == "| wide |  | end |"


def test_vertical_merge_not_duplicated():
    def fn(d):
        t = d.add_table(rows=3, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "k", "v"
        t.cell(1, 0).merge(t.cell(2, 0)).text = "group"
        t.cell(1, 1).text, t.cell(2, 1).text = "1", "2"

    t = parse(_build(fn), MIME).blocks[0]
    assert t.text.count("group") == 1


def test_lists_and_headers_footers_ignored():
    def fn(d):
        d.add_heading("Top", 1)
        d.add_paragraph("first", style="List Bullet")
        d.add_paragraph("second", style="List Number")
        d.add_paragraph("")
        d.add_paragraph("plain")
        d.sections[0].header.paragraphs[0].text = "HEADER TEXT"
        d.sections[0].footer.paragraphs[0].text = "FOOTER TEXT"

    p = parse(_build(fn), MIME)
    assert [b.type for b in p.blocks] == ["heading", "list_item", "list_item", "paragraph"]
    assert p.blocks[1].heading_path == ["Top"]
    assert "HEADER TEXT" not in p.text and "FOOTER TEXT" not in p.text


def test_canonical_offsets_corpus():
    p = _doc("travel-guidelines.docx")
    for b in p.blocks:
        assert p.text[b.char_start : b.char_end] == b.text


@pytest.mark.parametrize(
    "data",
    [b"not a zip at all", b"PK\x03\x04garbage", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100],
)
def test_corrupt_or_encrypted(data):
    with pytest.raises(IngestError) as e:
        parse(data, MIME)
    assert e.value.code == "INVALID_DOCUMENT"


def test_zip_that_is_not_docx():
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "hi")
    with pytest.raises(IngestError) as e:
        parse(buf.getvalue(), MIME)
    assert e.value.code == "INVALID_DOCUMENT"


def test_no_text_is_empty_extraction():
    with pytest.raises(IngestError) as e:
        parse(_build(lambda d: d.add_paragraph("   ")), MIME)
    assert e.value.code == "EMPTY_EXTRACTION"
